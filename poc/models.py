from django.db import models
from django.utils import timezone
from django.conf import settings
from openpyxl import Workbook
from openpyxl import load_workbook

import datetime

# Create your models here.

class Document(models.Model):
    path = models.FileField(upload_to='documents/')
    uploaded_at = models.DateTimeField(auto_now_add=True)
    description = models.CharField(max_length=100,default='none')
    test_set = models.IntegerField(default=0)
    remark = models.CharField(default="", max_length=200)

    def __str__(self):
        return self.description+'__'+str(self.test_set)
    def __open_collector(self,start_collect_row,key_index,filename,sheetname):
        values_dict = {}

        wb = load_workbook(filename = filename)
        worksheet = wb[sheetname]
        column_header_index = {}
        column_header = []
        ### header
        header_count = 0
        for row in worksheet['A'+str(start_collect_row):'Z'+str(start_collect_row)]:
            for cell in row:
                if cell.value != None:
                    column_header.append(cell.value)
                    column_header_index[cell.value] = header_count
                    header_count += 1
                else:
                    break
        ###
        index_cell = chr(ord('A') + column_header_index[key_index])
        start_cell_id = ord('A')
        for row in range(start_collect_row+1,1500):
            for col in range(0,25):
                iter = chr(start_cell_id+col)+str(row)
                if worksheet[iter].value != None:
                    fix_A_cell = index_cell+str(row)
                    if 'A' in iter:
                        values_dict[worksheet[fix_A_cell].value.strip()] = [worksheet[iter].value]
                    else:
                        values_dict[worksheet[fix_A_cell].value.strip()].append(worksheet[iter].value)
                else:
                    break
        return values_dict, column_header, column_header_index

    def __extrac_service_flow(self,flow_name):
        flow_name = flow_name.strip().upper()
        drop_time = {}
        if 'L2L3' in flow_name:
            service_type = 'L2L3'
        elif 'L3' in flow_name:
            service_type = 'L3'
        elif 'L2' in flow_name:
            service_type = 'L2'
        elif 'LOOP' in flow_name:
            service_type = 'LOOP'
        else:
            service_type = 'NONE'

        bg = 'true' if 'BG' in flow_name else 'false'

        sp = flow_name.split('_')
        if sp[2] > sp[3]:
            Aend = sp[2]
            Zend = sp[3]
            direction = 'downstream'
        elif sp[2] < sp[3]:
            Aend = sp[3]
            Zend = sp[2]
            direction = 'upstream'

        return service_type, bg, Aend, Zend, direction


    def upload_template(self,doctype):
        MEDIA_DIR = settings.BASE_DIR+str(self.path)
        insert_instances = []
        template_dict, template_header, theader_index = self.__open_collector(1,'Stream Block',MEDIA_DIR,'Result')
        for key,value in template_dict.items():
            try:
                flow = FlowTemplate.objects.get(flow_name=key.strip(), doctype=doctype)
                flow.fps = value[1]
                flow.id = str(self.id)
                flow.save()
            except FlowTemplate.DoesNotExist:
                insert_instances.append( FlowTemplate(flow_name=key.strip(), fps=value[1], document_id=str(self.id), doctype=doctype) )
        FlowTemplate.objects.bulk_create(insert_instances)

    def save_multicast_service_result(self, testcase, service_type, doctype):
        MEDIA_DIR = settings.BASE_DIR+str(self.path)
        template_dict, template_header, theader_index = self.__open_collector(4,'Stream Block',MEDIA_DIR,'Advanced Sequencing')
        summary = {}
        insert_instances = []
        mcast_srv_drop_list = [0]
        mcast_bg_drop_list = []
        time = timezone.now()
        if service_type == 'multicast_core':
            multiplier = 6
        elif service_type == 'multicast_headend':
            multiplier = 2   
        elif service_type == 'multicast_other':
            multiplier = 1
        # Flows
        for key,value in template_dict.items():
            try:
                fps = FlowTemplate.objects.get(flow_name=key.strip(), doctype=doctype).fps
            except FlowTemplate.DoesNotExist:
                continue
            bg = 'true' if 'BG' in key else 'false'            
            tx = value[theader_index['Tx Count (Frames)']]
            rx = value[theader_index['Rx Count (Frames)']]
            if 'LOOP' in key.strip().upper():
                drop_time = (tx - rx) / fps * 1000.0
                self.flows_set.create(  flow_name=key.strip(),
                        pub_date=time,
                        test_set=str(testcase),
                        tx=tx,
                        rx=rx,
                        drop_count=(tx - rx),
                        drop_time=round(drop_time,2),
                        service_type='LOOP',
                        bg_service=bg,
                        fps=fps)
            else:
                drop_time = ((6*tx - rx) / (multiplier * fps)) * 1000.0
                if bg == 'false':
                    mcast_srv_drop_list.append(drop_time)
                self.flows_set.create(  flow_name=key.strip(),
                                        pub_date=time,
                                        test_set=str(testcase),
                                        tx=tx,
                                        rx=rx,
                                        drop_count=(6*tx - rx),
                                        drop_time=round(drop_time,2),
                                        service_type=service_type,
                                        bg_service=bg,
                                        fps=fps)
        # FlowSummary Srv MCAST
        self.flowsummary_set.create(pub_date=time,
                                    test_set=str(testcase),
                                    A_end='Multicast',
                                    Z_end='Service',
                                    drop_time_upstream=0,
                                    drop_time_downstream=round(max(mcast_srv_drop_list),2),
                                    service_type=service_type,
                                    bg_service='false')

    def save_other_service_result(self,testcase, doctype):
        MEDIA_DIR = settings.BASE_DIR+str(self.path)
        template_dict, template_header, theader_index = self.__open_collector(4,'Stream Block',MEDIA_DIR,'Advanced Sequencing')
        summary = {}
        fps_dict = {}
        time = timezone.now()
        insert_instances = []
        # for flows
        query_list = FlowTemplate.objects.values('flow_name','fps','doctype')
        for row in query_list:
            if row['doctype'] not in fps_dict:
                fps_dict[row['doctype']] = {row['flow_name']: row['fps']}
            else:
                fps_dict[row['doctype']][row['flow_name']] = row['fps']

        for key,value in template_dict.items():
            ### check for flow dose not exist
            try:
                fps = fps_dict[doctype][key.strip()]
            except:
                continue
            ### check for tester error
            try:
                drop_count = value[theader_index['Tx-Rx (Frames)']]
                drop_time = drop_count / fps * 1000.0
            except:
                drop_count = -1
                drop_time = -1
            service_type, bg, Aend, Zend, direction = self.__extrac_service_flow(key)
            summary_key = Aend+'_'+Zend+'_'+bg+'_'+service_type

            if summary_key not in summary:
                if direction is 'upstream':
                    summary[summary_key] = {'Aend':Aend, 'Zend':Zend, 'upstream':drop_time, 'downstream':0, 'service_type':service_type, 'bg':bg}
                elif direction is 'downstream':
                    summary[summary_key] = {'Aend':Aend, 'Zend':Zend, 'upstream':0 ,'downstream':drop_time, 'service_type':service_type, 'bg':bg}
            else:
                if direction is 'upstream' and drop_time > summary[summary_key]['upstream']:
                    summary[summary_key]['upstream'] = drop_time
                elif direction is 'downstream' and drop_time > summary[summary_key]['downstream']:
                    summary[summary_key]['downstream'] = drop_time

            insert_instances.append(
                Flows(  flow_name=key.strip(),
                        pub_date=time,
                        test_set=str(testcase),
                        tx=value[theader_index['Tx Count (Frames)']],
                        rx=value[theader_index['Rx Count (Frames)']],
                        drop_count=drop_count,
                        drop_time=round(drop_time,2),
                        service_type=service_type,
                        bg_service=bg,
                        fps=fps,
                        document_id=self.id)
            )
        Flows.objects.bulk_create(insert_instances)
        ### for flowsummary
        insert_instances = []
        for k,v in summary.items():
            insert_instances.append(
                FlowSummary(    pub_date=time,
                                test_set=str(testcase),
                                A_end=v['Aend'],
                                Z_end=v['Zend'],
                                drop_time_upstream=round(v['upstream'],2),
                                drop_time_downstream=round(v['downstream'],2),
                                service_type=v['service_type'],
                                bg_service=v['bg'],
                                document_id=self.id)
            )
        FlowSummary.objects.bulk_create(insert_instances)

class FlowTemplate(models.Model):
    document = models.ForeignKey(Document, on_delete=models.CASCADE)
    flow_name = models.CharField(max_length=200)
    fps = models.IntegerField(default=1000)
    doctype = models.CharField(default="", max_length=100)
    def __str__(self):
        return self.flow_name

class Flows(models.Model):
    document = models.ForeignKey(Document, on_delete=models.CASCADE)
    flow_name = models.CharField(default="",max_length=200)
    pub_date = models.DateTimeField('date published')
    test_set = models.IntegerField(default=0)
    tx = models.BigIntegerField(default=0)
    rx = models.BigIntegerField(default=0)
    drop_count = models.BigIntegerField(default=0)
    drop_time = models.FloatField(default=0.0)
    service_type = models.CharField(max_length=100)
    bg_service = models.CharField(max_length=20,default='false')
    fps = models.IntegerField(default=0)

    def __str__(self):
        return self.flow_name+' drop: '+str(self.drop_count)

class FlowSummary(models.Model):
    document = models.ForeignKey(Document, on_delete=models.CASCADE)
    A_end = models.CharField(default="",max_length=100)
    Z_end = models.CharField(default="",max_length=100)
    test_set = models.IntegerField(default=0)
    service_type = models.CharField(max_length=100)
    bg_service = models.CharField(max_length=20,default='false')
    drop_time_upstream = models.FloatField(default=0.0)
    drop_time_downstream = models.FloatField(default=0.0)
    pub_date = models.DateTimeField('date published')
    def __str__(self):
        return self.A_end+' '+self.Z_end+' '+self.service_type+' '+self.bg_service

